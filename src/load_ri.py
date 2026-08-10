"""Load monthly Risk Index XLSX files (country × industry × month)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .utils import (
    PipelineError,
    excel_header_to_pandas,
    extract_month_from_filename,
    normalize_country_name,
    parse_ri_value,
)


def _list_ri_files(glob_path: Path) -> list[Path]:
    parent = glob_path.parent
    pattern = glob_path.name
    if not parent.exists():
        raise PipelineError(f"RI input directory missing: {parent}")
    files = sorted(
        p
        for p in parent.glob(pattern)
        if p.is_file() and not p.name.startswith("~$") and p.suffix.lower() == ".xlsx"
    )
    return files


def _validate_month_coverage(files: list[Path], year: int) -> dict[int, Path]:
    if len(files) != 12:
        raise PipelineError(
            f"Expected 12 RI monthly files, found {len(files)} under {files[0].parent if files else '?'}"
        )
    by_month: dict[int, Path] = {}
    for path in files:
        month = extract_month_from_filename(path, year)
        if month in by_month:
            raise PipelineError(
                f"Duplicate RI month {month}: {by_month[month].name} and {path.name}"
            )
        by_month[month] = path
    missing = [m for m in range(1, 13) if m not in by_month]
    if missing:
        raise PipelineError(f"Missing RI months: {missing}")
    return by_month


def _check_headers(actual: list[str], expected: list[str], path: Path) -> None:
    norm_actual = [str(c).strip() for c in actual]
    if norm_actual != expected:
        raise PipelineError(
            f"Unexpected RI headers in {path.name}: got {norm_actual}, expected {expected}"
        )


def _read_ri_sheet(path: Path, cfg_excel: dict[str, Any]) -> pd.DataFrame:
    sheet = cfg_excel["sheet_name"]
    header_row = excel_header_to_pandas(int(cfg_excel["excel_header_row"]))
    usecols = cfg_excel["usecols"]
    expected = cfg_excel["expected_headers"]

    # data_only=True → cached values only (no formula recalc / no external links)
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise PipelineError(f"Sheet {sheet!r} not found in {path.name}")
    wb.close()

    df = pd.read_excel(
        path,
        sheet_name=sheet,
        header=header_row,
        usecols=usecols,
        engine="openpyxl",
        dtype=str,
    )
    _check_headers(list(df.columns), expected, path)

    rename = {
        expected[0]: "country_name_raw",
        expected[1]: "industry_name",
        expected[2]: "country_code",
        expected[3]: "industry_code",
        expected[4]: "risk_index_raw",
    }
    df = df.rename(columns=rename)

    # Keep only rows with a Risk Index cell (actual data rows)
    df["risk_index_raw"] = df["risk_index_raw"].map(
        lambda x: "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x).strip()
    )
    df = df[df["risk_index_raw"] != ""].copy()
    if df.empty:
        raise PipelineError(f"No RI data rows in {path.name}")

    # Limited forward-fill for blank repeated country labels only
    for col in ("country_name_raw", "country_code"):
        blank = df[col].isna() | (df[col].astype(str).str.strip() == "") | (
            df[col].astype(str).str.lower() == "nan"
        )
        df.loc[blank, col] = pd.NA
        df[col] = df[col].ffill()

    df["country_name"] = df["country_name_raw"].map(normalize_country_name)
    df["country_code"] = df["country_code"].map(
        lambda x: "" if pd.isna(x) else str(x).strip()
    )
    df["industry_code"] = df["industry_code"].map(
        lambda x: "" if pd.isna(x) else str(x).strip()
    )
    df["industry_name"] = df["industry_name"].map(
        lambda x: "" if pd.isna(x) else str(x).strip()
    )

    if (df["country_name"] == "").any() or (df["country_code"] == "").any():
        raise PipelineError(
            f"Country name/code still blank after forward-fill in {path.name}"
        )

    try:
        df["ri"] = df["risk_index_raw"].map(parse_ri_value)
    except PipelineError as exc:
        raise PipelineError(f"{path.name}: {exc}") from exc

    if df["ri"].isna().any():
        raise PipelineError(f"Blank Risk Index after parse in {path.name}")

    return df[
        [
            "country_name",
            "country_code",
            "industry_name",
            "industry_code",
            "ri",
            "risk_index_raw",
        ]
    ]


def load_risk_index(cfg: dict[str, Any], logger) -> tuple[pd.DataFrame, dict]:
    from .utils import resolve_path

    year = int(cfg["analysis_year"])
    glob_path = resolve_path(cfg, cfg["paths"]["ri_glob"])
    files = _list_ri_files(glob_path)
    by_month = _validate_month_coverage(files, year)

    if not cfg.get("ri_higher_means_riskier", False):
        raise PipelineError(
            "Config ri_higher_means_riskier must be true for this analysis "
            "(RI 5 = high risk, RI 1 = low risk). Confirm report meta before flipping."
        )
    logger.info("RI direction confirmed: higher RI = higher risk (RI 5 worst)")

    frames = []
    for month, path in sorted(by_month.items()):
        part = _read_ri_sheet(path, cfg["ri_excel"])
        part["year"] = year
        part["month"] = month
        part["source_file"] = path.name
        frames.append(part)
        logger.info("Loaded RI %s month=%02d rows=%d", path.name, month, len(part))

    ri = pd.concat(frames, ignore_index=True)

    # Hard gate: one country_code must not map to many country names abnormally
    code_to_names = (
        ri.groupby("country_code")["country_name"].nunique().reset_index(name="n_names")
    )
    bad_codes = code_to_names[code_to_names["n_names"] > 1]
    if not bad_codes.empty:
        raise PipelineError(
            "Country code maps to multiple country names: "
            + ", ".join(
                f"{r.country_code}→{r.n_names}" for r in bad_codes.itertuples()
            )
        )

    profile = {
        "n_files": 12,
        "n_rows": int(len(ri)),
        "n_countries": int(ri["country_name"].nunique()),
        "n_country_codes": int(ri["country_code"].nunique()),
        "n_industries": int(ri["industry_code"].nunique()),
        "ri_value_counts": ri["ri"].value_counts().sort_index().to_dict(),
        "months": sorted(ri["month"].unique().tolist()),
    }
    return ri, profile
